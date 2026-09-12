"""Load structured dummy resumes into the shared candidate/evidence contract."""

from __future__ import annotations

import json
import argparse
import re
import xml.etree.ElementTree as ElementTree
from pathlib import Path
from typing import Any

from src.matching.normalization import extract_skill_mentions
from src.parsing.section_detector import detect_sections
from src.parsing.text_normalizer import normalize_text, normalize_value


def load_resume_directory(directory: Path, extensions: tuple[str, ...] = ("xml",)) -> list[dict[str, Any]]:
    files = [path for extension in extensions for path in sorted(directory.glob(f"*.{extension}"))]
    return [load_resume(path) for path in files]


def load_resume(path: Path) -> dict[str, Any]:
    if path.suffix.lower() == ".xml":
        return load_xml_resume(path)
    if path.suffix.lower() == ".pdf":
        return load_pdf_resume(path)
    if path.suffix.lower() == ".docx":
        return load_docx_resume(path)
    if path.suffix.lower() == ".txt":
        return load_text_resume(path)
    raise ValueError(f"Unsupported resume format: {path.suffix}")


def load_pdf_resume(path: Path) -> dict[str, Any]:
    try:
        import fitz
    except ImportError as error:
        raise RuntimeError("PDF support requires pymupdf") from error

    document = fitz.open(path)
    page_text = [normalize_text(page.get_text("text")) for page in document]
    if needs_ocr(page_text):
        page_text = ocr_pdf_pages(document)
    raw_text = normalize_text("\n".join(page_text))
    return build_candidate(
        path,
        first_line(raw_text, path.stem),
        detect_sections(raw_text),
        raw_text,
        source_pages=page_text,
    )


def needs_ocr(page_text: list[str], minimum_characters: int = 40) -> bool:
    if not page_text:
        return True
    low_text_pages = sum(len(text.strip()) < minimum_characters for text in page_text)
    return low_text_pages > len(page_text) / 2


def ocr_pdf_pages(document: Any) -> list[str]:
    try:
        import fitz
        import pytesseract
        from PIL import Image
    except ImportError as error:
        raise RuntimeError("OCR support requires pytesseract and pillow") from error

    pages = []
    for page in document:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        pages.append(normalize_text(pytesseract.image_to_string(image)))
    return pages


def load_docx_resume(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError("DOCX support requires python-docx") from error

    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_values = [cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
    values = paragraphs + table_values
    raw_text = normalize_text("\n".join(values))
    return build_candidate(path, values[0] if values else path.stem, detect_sections(raw_text), raw_text)


def load_xml_resume(path: Path) -> dict[str, Any]:
    root = ElementTree.parse(path).getroot()
    applicant = root.find("applicant")
    candidate_name = text_at(applicant, "name") if applicant is not None else path.stem
    sections: dict[str, Any] = {
        "summary": text_at(root, "summary"),
        "skills": [text_value(node) for node in root.findall("./technicalSkills/skillGroup")],
        "experience": section_items(root, "experience"),
        "projects": project_items(root),
        "education": section_items(root, "education"),
        "certifications": section_items(root, "certifications"),
    }
    raw_text = normalize_text("\n\n".join(value for value in flatten_sections(sections) if value))
    return build_candidate(path, candidate_name, sections, raw_text)


def load_text_resume(path: Path) -> dict[str, Any]:
    raw_text = normalize_text(path.read_text(encoding="utf-8"))
    candidate_name = first_line(raw_text, path.stem)
    sections = detect_sections(raw_text)
    return build_candidate(path, candidate_name, sections, raw_text)


def first_line(value: str, fallback: str) -> str:
    return next((line.strip() for line in value.splitlines() if line.strip()), fallback)


def build_candidate(
    path: Path,
    candidate_name: str,
    sections: dict[str, Any],
    raw_text: str,
    source_pages: list[str] | None = None,
) -> dict[str, Any]:
    candidate_id = slugify(path.stem)
    evidence = []
    position = 1
    for section, values in sections.items():
        for value in flatten_value(values):
            value = normalize_value(value)
            if not value:
                continue
            skills = find_skills(value)
            evidence.append(
                {
                    "candidate_id": candidate_id,
                    "evidence_id": f"{candidate_id}_evidence_{position:03d}",
                    "section": section,
                    "text": value,
                    "page": find_page(value, source_pages),
                    "position": position,
                    "extracted_skill": skills[0] if skills else None,
                    "canonical_skill": skills[0] if skills else None,
                    "evidence_type": "skill_list" if section == "skills" else "project_or_experience",
                    "source": section,
                    "evidence_source_type": "skills_section" if section == "skills" else section,
                    "evidence_strength": "applied" if section == "skills" else "substantial",
                }
            )
            position += 1

    return {
        "candidate_id": candidate_id,
        "candidate_name": candidate_name,
        "resume_source": path.name,
        "raw_text": raw_text,
        "sections": sections,
        "evidence": evidence,
        "metadata": {"source_format": path.suffix.lower().lstrip(".")},
    }


def find_page(value: str, source_pages: list[str] | None) -> int | None:
    if not source_pages:
        return None
    normalized_value = normalize_value(value).lower()
    for page_number, page_text in enumerate(source_pages, start=1):
        if normalized_value and normalized_value in page_text.lower():
            return page_number
    return None


def section_items(root: ElementTree.Element, section_name: str) -> list[str]:
    section = root.find(section_name)
    if section is None:
        return []
    return [" ".join(part.strip() for part in element.itertext() if part.strip()) for element in section]


def project_items(root: ElementTree.Element) -> list[str]:
    values = []
    for project in root.findall("./projects/project"):
        values.append(" ".join(part.strip() for part in project.itertext() if part.strip()))
    return values


def text_at(parent: ElementTree.Element | None, child_name: str) -> str:
    if parent is None:
        return ""
    child = parent.find(child_name)
    return text_value(child) if child is not None else ""


def text_value(element: ElementTree.Element | None) -> str:
    return " ".join(element.itertext()).strip() if element is not None else ""


def flatten_sections(sections: dict[str, Any]) -> list[str]:
    values = []
    for section_values in sections.values():
        values.extend(flatten_value(section_values))
    return values


def flatten_value(value: Any) -> list[str]:
    if isinstance(value, list):
        return [item for item in value if isinstance(item, str)]
    if isinstance(value, str):
        return [value]
    return []


def find_skills(text: str) -> list[str]:
    return extract_skill_mentions(text)


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--extensions",
        nargs="+",
        default=["xml"],
        choices=["pdf", "docx", "txt", "xml"],
        help="Resume formats to load; defaults to the structured XML fixture.",
    )
    args = parser.parse_args()
    project_root = Path(__file__).resolve().parents[2]
    source = project_root / "data" / "dummy_resumes" / "Dummy Resumes"
    output = project_root / "data" / "dummy_candidates.json"
    candidates = load_resume_directory(source, tuple(args.extensions))
    output.write_text(json.dumps({"candidates": candidates}, indent=2) + "\n", encoding="utf-8")
    print(f"Loaded {len(candidates)} {', '.join(args.extensions)} resumes into {output}")
